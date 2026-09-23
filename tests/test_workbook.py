"""Workbook: ES duplo/independente, persistencia em copia temporaria + backup.

Nunca usa o DADOS.xlsx real em teste destrutivo: copia para tmp_path.
"""
import shutil
import unicodedata

import openpyxl
import pytest

from src import config
from src.workbook_service import parse_entities


def _copy_real_to(tmp_path, monkeypatch):
    dest = tmp_path / "DADOS.xlsx"
    shutil.copy2(str(config.WORKBOOK_PATH), str(dest))
    monkeypatch.setattr(config, "WORKBOOK_PATH", dest)
    import src.workbook_service as wbs
    monkeypatch.setattr(wbs, "WORKBOOK_PATH", dest)
    monkeypatch.setattr(wbs, "BACKUPS_DIR", tmp_path / "BACKUPS")
    monkeypatch.setattr(wbs, "LOCK_PATH", tmp_path / ".dados.lock")
    return dest


def test_es_duplo_independente(tmp_path, monkeypatch):
    dest = _copy_real_to(tmp_path, monkeypatch)
    import src.workbook_service as wbs
    wb = openpyxl.load_workbook(str(dest))
    try:
        ents = parse_entities(wb)
        keys = [e["entity_key"] for e in ents]
        assert "ES_PP" in keys and "ES_SEJUS" in keys
        assert len(ents) == 28
        pp = next(e for e in ents if e["entity_key"] == "ES_PP")
        sej = next(e for e in ents if e["entity_key"] == "ES_SEJUS")
        assert pp["row"] != sej["row"]
        assert pp["entity_key"] != sej["entity_key"]
        assert pp["unidade_label"] != sej["unidade_label"]
    finally:
        wb.close()


def test_persistencia_status_backup(tmp_path, monkeypatch):
    dest = _copy_real_to(tmp_path, monkeypatch)
    import src.workbook_service as wbs
    res = wbs.update_assessment("AC", "02_Autonomia:M3-56", status="Atende")
    assert res["ok"] is True
    assert res["score"] == 3
    backups = list((tmp_path / "BACKUPS").glob("DADOS_*.xlsx"))
    assert len(backups) >= 1
    wb = openpyxl.load_workbook(str(dest))
    try:
        ws = wb["02_Autonomia"]
        assert ws.cell(7, 3).value == "Atende"
        # abas auxiliares criadas
        assert "DB_EVIDENCIAS" in wb.sheetnames
        assert "DB_ANEXOS" in wb.sheetnames
        assert "DB_AUDITORIA" in wb.sheetnames
        # evidencia inicial
        wse = wb["DB_EVIDENCIAS"]
        texts = [wse.cell(r, 10).value for r in range(2, wse.max_row + 1)]
        assert config.EVIDENCE_INITIAL_TEXT in texts
    finally:
        wb.close()
    # reabrir confirma persistencia
    wb2 = openpyxl.load_workbook(str(dest))
    assert wb2["02_Autonomia"].cell(7, 3).value == "Atende"
    wb2.close()


def test_save_atomic_retries_and_cleans_temp_when_workbook_is_locked(tmp_path, monkeypatch):
    dest = _copy_real_to(tmp_path, monkeypatch)
    import src.workbook_service as wbs

    wb = wbs.open_workbook()
    attempts = []

    def fail_replace(_source, _target):
        attempts.append(True)
        raise PermissionError(32, "The file is being used by another process")

    monkeypatch.setattr(wbs.Path, "replace", fail_replace)
    monkeypatch.setattr("time.sleep", lambda _delay: None)
    try:
        with pytest.raises(wbs.LockedError, match="Feche a planilha no Excel"):
            wbs.save_atomic(wb)
        assert len(attempts) == 6
        assert not list(tmp_path.glob("DADOS_tmp_*.xlsx"))
        assert dest.exists()
    finally:
        wb.close()


def test_api_patch_assessment_flow(tmp_path, monkeypatch):
    dest = _copy_real_to(tmp_path, monkeypatch)
    from urllib.parse import quote
    from starlette.testclient import TestClient
    import app as app_module
    import src.workbook_service as wbs

    # O arquivo pode guardar o nome da aba em NFD enquanto o frontend envia NFC.
    wb = openpyxl.load_workbook(str(dest))
    try:
        canonical = "01_Institucionalização"
        ws = next(s for s in wb.worksheets if unicodedata.normalize("NFC", s.title) == canonical)
        ws.title = unicodedata.normalize("NFD", canonical)
        wb.save(str(dest))
    finally:
        wb.close()

    monkeypatch.setattr(app_module, "STARTUP_ERRORS", [])
    client = TestClient(app_module.app)

    # 1. Sem status e sem evidencia -> 422 com mensagem clara
    occ = quote("01_Institucionalização:M1-11")
    resp_empty = client.patch(f"/api/unidades/AC/avaliacoes/{occ}", json={"status": None, "evidence_text": None})
    assert resp_empty.status_code == 422
    assert "Informe status e/ou evidence_text" in resp_empty.json()["detail"]

    # 2. Com status "Não" e evidencia
    resp = client.patch(f"/api/unidades/AC/avaliacoes/{occ}", json={"status": "Não", "evidence_text": "Comprovante de teste"})
    assert resp.status_code == 200
    data = resp.json()
    assert data["ok"] is True
    assert data["score"] == 0.0
    assert data["classification"] == "Não instituída"
    assert data["situacao"] == "Não instituída"

    # A API grava na célula da pergunta e preserva a estrutura do XLSX.
    wb = openpyxl.load_workbook(str(dest))
    try:
        ent = next(e for e in wbs.parse_entities(wb) if e["entity_key"] == "AC")
        q = wbs.find_question(wb, "01_Institucionalização:M1-11")
        ws = wbs.get_sheet_by_name(wb, q["sheet_name"])
        assert ws.cell(ent["row"], q["status_col"]).value == "Não"
        evidence_col = wbs.find_evidence_col(ws)
        assert evidence_col and ws.cell(ent["row"], evidence_col).value == "Comprovante de teste"
        evidence = wb["DB_EVIDENCIAS"]
        assert any(
            evidence.cell(row, 2).value == "AC"
            and evidence.cell(row, 9).value == "01_Institucionalização:M1-11"
            and evidence.cell(row, 10).value == "Comprovante de teste"
            for row in range(2, evidence.max_row + 1)
        )
        audit = wb["DB_AUDITORIA"]
        assert any(
            audit.cell(row, 3).value == "AC"
            and audit.cell(row, 6).value == "01_Institucionalização:M1-11"
            and audit.cell(row, 8).value == "Sim"
            and audit.cell(row, 9).value == "Não"
            for row in range(2, audit.max_row + 1)
        )
    finally:
        wb.close()


def test_api_patch_returns_conflict_when_workbook_is_locked(monkeypatch):
    from urllib.parse import quote
    from starlette.testclient import TestClient
    import app as app_module
    from src.workbook_service import LOCKED_MSG

    monkeypatch.setattr(app_module, "STARTUP_ERRORS", [LOCKED_MSG])
    monkeypatch.setattr(app_module, "validate_startup", lambda: [LOCKED_MSG])
    client = TestClient(app_module.app)
    occurrence = quote("01_Institucionalização:M1-11")

    response = client.patch(
        f"/api/unidades/AC/avaliacoes/{occurrence}",
        json={"status": "Não", "evidence_text": None},
    )

    assert response.status_code == 409
    assert "Feche a planilha no Excel e tente novamente" in response.json()["detail"]
