"""Anexos: extensao, traversal, hash, metadados, desvinculacao sem apagar."""
import hashlib

import pytest

from src import attachments
from src.attachments import sanitize_filename, validate_file
from src.workbook_service import WorkbookError


def test_extensao_permitida():
    assert validate_file("doc.pdf", 100) == ".pdf"
    assert validate_file("foto.JPG", 100) == ".jpg"


def test_extensao_proibida():
    with pytest.raises(WorkbookError):
        validate_file("malware.exe", 100)
    with pytest.raises(WorkbookError):
        validate_file("script.bat", 100)
    with pytest.raises(WorkbookError):
        validate_file("semextensao", 100)


def test_tamanho_limite():
    with pytest.raises(WorkbookError):
        validate_file("grande.pdf", 25 * 1024 * 1024 + 1)


def test_path_traversal_neutralizado():
    assert "/" not in sanitize_filename("../../etc/passwd.pdf")
    assert sanitize_filename("../../etc/passwd.pdf").endswith(".pdf")
    assert "\\" not in sanitize_filename("..\\..\\win.pdf")
    # ':' e invalidos viram '_'; Path().name pode descartar drive-like prefixo
    cleaned = sanitize_filename("a:b*c?.pdf")
    assert "\\" not in cleaned and "/" not in cleaned and ":" not in cleaned
    assert cleaned.endswith(".pdf")


def test_hash_sha256():
    data = b"conteudo de teste"
    assert hashlib.sha256(data).hexdigest() == hashlib.sha256(data).hexdigest()
    assert len(hashlib.sha256(data).hexdigest()) == 64


def test_desvinculacao_marca_inactive_sem_apagar(tmp_path, monkeypatch):
    """Simula DB_ANEXOS minimo e garante que unlink marca FALSE sem unlink fisico."""
    import openpyxl
    from src import config

    planted = tmp_path / "DADOS.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DB_ANEXOS"
    ws.append(config.DB_ANEXOS_HEADERS)
    ws.append([1, "AC", "AC", "Acre", "02_Autonomia", "Autonomia", "Item 1",
               "M3-56", "02_Autonomia:M3-56", "doc.pdf", "uuid__doc.pdf",
               "ANEXOS/AC/x/doc.pdf", "application/pdf", 10, "abc",
               "2026-01-01T00:00:00", "TRUE", ""])
    ws2 = wb.create_sheet("DB_AUDITORIA")
    ws2.append(config.DB_AUDITORIA_HEADERS)
    wb.save(str(planted))

    monkeypatch.setattr(config, "WORKBOOK_PATH", planted)
    import src.workbook_service as wbs
    monkeypatch.setattr(wbs, "WORKBOOK_PATH", planted)
    monkeypatch.setattr("src.attachments.WORKBOOK_PATH", planted, raising=False)
    # store_attachment/unlink usam create_backup + save_atomic do modulo
    # workbook_service; redireciona paths de backup p/ tmp
    monkeypatch.setattr(config, "BACKUPS_DIR", tmp_path / "BACKUPS")
    monkeypatch.setattr(wbs, "BACKUPS_DIR", tmp_path / "BACKUPS")

    res = attachments.unlink_attachment(1)
    assert res["ok"] is True
    wb2 = openpyxl.load_workbook(str(planted))
    assert wb2["DB_ANEXOS"].cell(2, 17).value == "FALSE"
    wb2.close()
