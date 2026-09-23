"""Relatorios PDF (ReportLab) e XLSX (openpyxl) — gerais e individuais."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

from .config import (
    BASE_WEIGHTS,
    DIMENSION_MINIMUMS,
    DIMENSION_NAMES,
    EXPORTACOES_DIR,
    GLOBAL_BASE_MINIMUM,
)
from .workbook_service import read_unit_detail

HEADER_LINES = [
    "MINISTÉRIO DA JUSTIÇA E SEGURANÇA PÚBLICA",
    "SECRETARIA NACIONAL DE POLÍTICAS PENAIS",
    "OUVIDORIA NACIONAL DE SERVIÇOS PENAIS",
    "MONITORAMENTO DOS PARÂMETROS MÍNIMOS DAS OUVIDORIAS DE SERVIÇOS PENAIS",
]
NOTA_METODOLOGICA = (
    "As pontuações, pesos e faixas de classificação constituem metodologia "
    "de monitoramento da ONASP e não integram o texto da Instrução Normativa "
    "GABSEC/SENAPPEN/MJSP nº 75/2026. A avaliação deve ser interpretada em "
    "conjunto com as evidências registradas no sistema."
)


LOGO_PATH = Path(__file__).resolve().parent.parent / "static" / "logos" / "logo_senappen_mjsp_gov_horizontal.png"
SEI_PROCESSO_URL = "https://sei.mj.gov.br/sei/controlador.php?acao=procedimento_trabalhar&amp;id_procedimento=38337209"
IN_75_URL = "https://sei.mj.gov.br/sei/controlador.php?acao=procedimento_trabalhar&amp;id_procedimento=38337209&amp;id_documento=39918839"


def _stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def _base_doc(path: Path, title: str) -> tuple[SimpleDocTemplate, list, dict]:
    EXPORTACOES_DIR.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(path), pagesize=A4, title=title,
                            leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm)
    styles = getSampleStyleSheet()
    return doc, [], styles


def _header_flow(styles, report_subtitle: str = "") -> list:
    header_org_style = ParagraphStyle(
        "HeaderOrg",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#103E49"),
    )
    header_title_style = ParagraphStyle(
        "HeaderSystemTitle",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#17365D"),
    )
    header_sub_style = ParagraphStyle(
        "HeaderReportSubtitle",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=13.5,
        textColor=colors.HexColor("#103E49"),
    )
    header_ref_style = ParagraphStyle(
        "HeaderRef",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor("#526572"),
    )

    flows = []
    if LOGO_PATH.exists():
        logo_w = 90 * mm
        logo_h = logo_w * (197 / 1342)  # proporção exata 1342 x 197 (~13.2mm)
        img = Image(str(LOGO_PATH), width=logo_w, height=logo_h)
        img.hAlign = "CENTER"
        flows.append(img)
        flows.append(Spacer(1, 2 * mm))
    else:
        for line in HEADER_LINES[:2]:
            flows.append(Paragraph(line, header_org_style))
        flows.append(Spacer(1, 1.5 * mm))

    # Ouvidoria Nacional de Serviços Penais
    flows.append(Paragraph(HEADER_LINES[2], header_org_style))
    flows.append(Spacer(1, 1.5 * mm))
    # Nome do Sistema centralizado
    flows.append(Paragraph(HEADER_LINES[3], header_title_style))
    if report_subtitle:
        flows.append(Spacer(1, 1.5 * mm))
        flows.append(Paragraph(report_subtitle, header_sub_style))
    flows.append(Spacer(1, 1.5 * mm))
    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    flows.append(Paragraph(
        f"Referência: IN GABSEC/SENAPPEN/MJSP nº 75/2026 — metodologia ONASP. "
        f"Gerado em {now}.", header_ref_style))
    flows.append(Spacer(1, 3.5 * mm))
    return flows


def _methodology_flow(styles) -> list:
    annex_header_style = ParagraphStyle(
        "AnnexHeader",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=13.5,
        textColor=colors.HexColor("#103E49"),
    )
    annex_sub_style = ParagraphStyle(
        "AnnexSub",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#17365D"),
    )
    annex_ref_style = ParagraphStyle(
        "AnnexRef",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontSize=7,
        leading=9.5,
        textColor=colors.HexColor("#526572"),
    )
    body_style = ParagraphStyle(
        "AnnexBody",
        parent=styles["Normal"],
        fontSize=7,
        leading=9.5,
        alignment=TA_JUSTIFY,
        textColor=colors.HexColor("#192E39"),
    )
    h2_style = ParagraphStyle(
        "AnnexH2",
        parent=styles["Heading2"],
        fontSize=8.5,
        leading=11.5,
        fontName="Helvetica-Bold",
        textColor=colors.HexColor("#103E49"),
        spaceBefore=2.5 * mm,
        spaceAfter=1 * mm,
    )
    note_box_style = ParagraphStyle(
        "AnnexNote",
        parent=styles["Normal"],
        fontSize=6.5,
        leading=9,
        alignment=TA_JUSTIFY,
        textColor=colors.HexColor("#103E49"),
    )

    flows = [PageBreak()]
    if LOGO_PATH.exists():
        logo_w = 60 * mm
        logo_h = logo_w * (197 / 1342)  # ~8.8mm
        img = Image(str(LOGO_PATH), width=logo_w, height=logo_h)
        img.hAlign = "CENTER"
        flows.append(img)
        flows.append(Spacer(1, 1.5 * mm))

    flows.extend([
        Paragraph("<b>ANEXO</b>", annex_sub_style),
        Spacer(1, 0.8 * mm),
        Paragraph("<b>METODOLOGIA DE MONITORAMENTO DOS PARÂMETROS MÍNIMOS DAS OUVIDORIAS DE SERVIÇOS PENAIS</b>", annex_header_style),
        Spacer(1, 0.8 * mm),
        Paragraph(
            f"<b>Referência:</b> Processo <a href='{SEI_PROCESSO_URL}' color='#155b67'><u>SEI nº 08016.027689/2025-19</u></a> · "
            f"<a href='{IN_75_URL}' color='#155b67'><u>IN GABSEC/SENAPPEN/MJSP nº 75/2026</u></a> · Doc. SEI 37070578",
            annex_ref_style
        ),
        Spacer(1, 2 * mm),

        Paragraph("<b>1. Apresentação e Finalidade</b>", h2_style),
        Paragraph(
            "Este sistema foi desenvolvido pela Ouvidoria Nacional de Serviços Penais (ONASP/SENAPPEN/MJSP) para apoiar o "
            "monitoramento da institucionalização, da estruturação e das condições de funcionamento das Ouvidorias de Serviços "
            "Penais dos Estados e do Distrito Federal. A ferramenta articula três componentes principais: as metas do Plano Nacional "
            "Pena Justa; os parâmetros mínimos da Instrução Normativa GABSEC/SENAPPEN/MJSP nº 75/2026; e as informações fáticas do "
            "diagnóstico nacional realizado pela ONASP junto às Unidades Federativas. Sua finalidade é organizar evidências, permitir "
            "comparação padronizada, acompanhar a evolução da política pública e orientar ações federais de fomento institucional.", body_style),

        Paragraph("<b>2. Relação com o Plano Nacional Pena Justa</b>", h2_style),
        Paragraph(
            "O monitoramento atende diretamente às metas de implementação do Plano Nacional Pena Justa, em especial: "
            "<b>Indicador 2.4.2.1.1.1</b> (Elaboração de parâmetros para criação de ouvidorias estaduais autônomas de serviços penais) e "
            "<b>Indicador 2.4.2.1.2.1</b> (Estabelecimento de ouvidorias estaduais criadas, seguindo os parâmetros mínimos). "
            "Primeiro estabeleceu-se a referência normativa (IN nº 75/2026); em seguida afere-se a aderência concreta das unidades.", body_style),

        Paragraph("<b>3. Estrutura da Avaliação, Dimensões e Pesos</b>", h2_style),
        Paragraph(
            "A matriz de avaliação é composta por 6 dimensões de pontuação-base (totalizando até 100 pontos) e 1 bloco de maturidade/bônus "
            "(até 10 pontos compensatórios). Para assegurar rigor e evitar dupla contagem, cada pergunta do diagnóstico vincula-se "
            "exclusivamente a uma única dimensão:", body_style),
        Spacer(1, 1 * mm),
    ])

    weights_table_data = [
        ["Dimensão", "Peso Máx.", "Piso Mínimo Exigido (IN nº 75/2026)"],
        ["01. Institucionalização", "15 pts", "Requisito de entrada: ato formal publicado (binário: 15 ou 0)"],
        ["02. Autonomia técnica e funcional", "15 pts", "Piso de 50% = 7,5 pontos"],
        ["03. Imparcialidade, sigilo e proteção", "15 pts", "Piso de 50% = 7,5 pontos"],
        ["04. Acessibilidade e atendimento humanizado", "15 pts", "Piso de 50% = 7,5 pontos"],
        ["05. Transparência e publicidade", "15 pts", "Piso de 50% = 7,5 pontos"],
        ["06. Integração tecnológica", "25 pts", "Piso de 50% = 12,5 pontos"],
        ["Total da Nota-base", "100 pts", "Piso global: 70 pontos na Nota Final"],
        ["07. Maturidade / Bônus adicional", "10 pts bônus", "Compensatório: preenche a nota até 100 e apoia a nota global"],
    ]
    tw = Table(weights_table_data, colWidths=[60 * mm, 28 * mm, 92 * mm])
    tw.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("ALIGN", (1, 1), (1, -1), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 7), (-1, 7), "Helvetica-Bold"),
        ("BACKGROUND", (0, 7), (-1, 7), colors.HexColor("#EAF2F4")),
    ]))
    flows.append(tw)
    flows.append(Spacer(1, 1.5 * mm))

    flows.extend([
        Paragraph("<b>4. Critérios de Pontuação das Perguntas</b>", h2_style),
        Paragraph(
            "Cada pergunta ordinária é pontuada conforme o nível de comprovação registrado: "
            "<b>Atende (100% dos pontos):</b> evidência demonstra atendimento integral ao critério; "
            "<b>Parcial (50% dos pontos):</b> atendimento intermediário ou em implementação; "
            "<b>Não atende (0%):</b> evidência aponta ausência do requisito; "
            "<b>Sem evidência (0%):</b> ausência de documentação ou comprovação suficiente no momento da avaliação.", body_style),

        Paragraph("<b>5. Regras para Cumprimento dos Parâmetros Mínimos e Faixas de Classificação</b>", h2_style),
        Paragraph(
            "Uma unidade avaliada é considerada <b>“Instituída — seguindo os parâmetros mínimos”</b> quando cumpre "
            "cumulativamente três condições essenciais: (1) Possui ato normativo instituidor (M1-11 = Sim / Instituída); "
            "(2) Atinge nota final igual ou superior a 70 pontos (nota-base somada ao bônus aplicado); e "
            "(3) Atende ao piso mínimo regulamentar de 50% em cada uma das dimensões essenciais 02, 03, 04, 05 e 06.", body_style),
        Spacer(1, 1 * mm),
    ])

    cls_table_data = [
        ["Classificação Metodológica", "Regra Objetiva / Condições"],
        ["Não instituída", "Não possui ato normativo de criação (M1-11 = Não)."],
        ["Instituição não comprovada", "Ausência de comprovação do ato normativo de criação (M1-11 = Sem evidência)."],
        ["Instituída — abaixo do mínimo em dimensão essencial", "Possui ato normativo, mas não atingiu 50% em uma ou mais dimensões essenciais (02 a 06)."],
        ["Instituída — aderência global insuficiente", "Possui ato normativo e atendeu aos pisos dimensionais, mas a nota final ficou abaixo de 70 pontos."],
        ["Instituída — seguindo os parâmetros mínimos", "Possui ato normativo, atendeu a todos os pisos dimensionais (≥ 50%) e alcançou nota final ≥ 70 pontos."],
    ]
    tc = Table(cls_table_data, colWidths=[68 * mm, 112 * mm])
    tc.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
    ]))
    flows.append(tc)
    flows.append(Spacer(1, 1.5 * mm))

    flows.extend([
        Paragraph("<b>6. Rastreabilidade Documental e Limites da Metodologia</b>", h2_style),
        Paragraph(
            "Toda a análise baseia-se nas informações prestadas pelas Unidades Federativas no diagnóstico nacional (referência geral "
            "<b>Doc. SEI 37070578</b>) e nos atos normativos, relatórios e documentos comprobatórios anexados a cada pergunta. "
            "As pontuações, pesos e faixas de classificação constituem instrumento técnico e gerencial de monitoramento da ONASP/SENAPPEN "
            "e não integram textualmente a Instrução Normativa nº 75/2026. A pontuação não constitui sanção, certificação jurídica "
            "nem substitui a análise jurídica individualizada de cada ato normativo estadual.", body_style),
        Spacer(1, 2.5 * mm),
    ])

    box_data = [[
        Paragraph(
            "<b>AVISO INSTITUCIONAL:</b> As pontuações, pesos e faixas de classificação constituem metodologia "
            "de monitoramento da ONASP e não integram o texto da Instrução Normativa GABSEC/SENAPPEN/MJSP nº 75/2026. "
            "A avaliação deve ser interpretada em conjunto com as evidências registradas no sistema.", note_box_style)
    ]]
    tbox = Table(box_data, colWidths=[180 * mm])
    tbox.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F4F6F7")),
        ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor("#103E49")),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    flows.append(tbox)
    return flows


def _table_style() -> TableStyle:
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),  # cabeçalhos da tabela centralizados
    ])


def generate_general_pdf(rows: list[dict], cards: dict) -> Path:
    path = EXPORTACOES_DIR / f"relatorio_geral_{_stamp()}.pdf"
    doc, story, styles = _base_doc(path, "Relatório geral — Parâmetros Mínimos")
    story.extend(_header_flow(styles, "Relatório Geral Consolidado"))
    story.append(Paragraph("<b>Síntese</b>", styles["Heading2"]))
    story.append(Paragraph(
        f"Unidades avaliadas: {cards['unidades']}. Ouvidorias instituídas: {cards['instituidas']}. "
        f"Não instituídas: {cards['nao_instituidas']}. Instituição não comprovada: {cards['nao_comprovadas']}. "
        f"Seguindo os parâmetros mínimos: {cards['seguindo_minimos']}. "
        f"Abaixo do mínimo em dimensão essencial: {cards['abaixo_dimensao']}. "
        f"Aderência global insuficiente: {cards['global_insuficiente']}.",
        styles["Normal"]))
    story.append(Spacer(1, 4 * mm))
    data = [["UF", "Unidade", "Inst", "Aut", "Imp", "Aces", "Trans", "Integ",
             "Base", "Bônus", "Final", "Classificação"]]
    for r in rows:
        d = r["dim"]
        data.append([
            r["uf"], r["unidade_label"][:40],
            _fmt(d.get("01_Institucionalização")), _fmt(d.get("02_Autonomia")),
            _fmt(d.get("03_Imparcialidade")), _fmt(d.get("04_Acessibilidade")),
            _fmt(d.get("05_Transparência")), _fmt(d.get("06_Integração Tec")),
            _fmt(r["base_score"]), _fmt(r["bonus_applied"]), _fmt(r["final_score"]),
            (r["classification"] or "")[:60],
        ])
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),  # cabeçalhos centralizados
        ("ALIGN", (0, 1), (0, -1), "CENTER"),  # coluna UF centralizada
        ("ALIGN", (2, 1), (-2, -1), "CENTER"),  # pontuações centralizadas
    ]))
    story.append(t)
    story.extend(_methodology_flow(styles))
    doc.build(story)
    return path


def generate_unit_pdf(entity_key: str) -> Path:
    detail = read_unit_detail(entity_key)
    ent = detail["entity"]
    res = detail["result"]
    path = EXPORTACOES_DIR / f"relatorio_{entity_key}_{_stamp()}.pdf"
    doc, story, styles = _base_doc(path, f"Relatório individual — {entity_key}")
    story.extend(_header_flow(styles, f"Relatório de Avaliação Individual — {ent['uf']} — {ent['unidade_label']}"))
    unit_meta_style = ParagraphStyle(
        "UnitMeta",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#192E39"),
    )
    story.append(Paragraph(
        f"<b>UF:</b> {ent['uf']} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Unidade:</b> {ent['unidade_label']} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Situação:</b> {res['situacao']} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Classificação:</b> {res['classification']}", unit_meta_style))
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph("<b>Resultado Consolidado</b>", styles["Heading2"]))
    dimension_results = " | ".join(
        f"{dim['dimension_name']}: {_fmt(res['dim_scores'].get(dim['sheet_name']))}"
        + (f" (mínimo {DIMENSION_MINIMUMS[dim['sheet_name']]:g})"
           if dim["sheet_name"] in DIMENSION_MINIMUMS else "")
        for dim in detail["dimensions"]
        if dim["sheet_name"] in BASE_WEIGHTS
    )
    story.append(Paragraph(
        f"{dimension_results} | "
        f"Nota-base: {_fmt(res['base_score'])} | Bônus disponível: {_fmt(res['bonus_available'])} | "
        f"Bônus aplicado: {_fmt(res['bonus_applied'])} | Nota final (mínimo global {GLOBAL_BASE_MINIMUM}): {_fmt(res['final_score'])} | "
        f"Segue os parâmetros mínimos: {'Sim' if res['meets_minimum_parameters'] else 'Não'} | "
        f"Classificação: {res['classification']}", styles["Normal"]))
    for dim in detail["dimensions"]:
        story.append(Paragraph(f"<b>{dim['dimension_name']}</b>", styles["Heading3"]))
        qdata = []
        for q in dim["questions"]:
            docs = ", ".join(a["original_filename"] for a in q["attachments"]) or "—"
            raw_fund = (q.get("fundamentacao") or "").strip()
            if raw_fund:
                escaped_fund = raw_fund[:200].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                fund_link = f"<a href='{IN_75_URL}' color='#155b67'><u>{escaped_fund}</u></a>"
            else:
                fund_link = "—"
            teor_plain = (q.get("fundamentacao_teor_plain") or "").strip()
            teor_flow = ""
            if teor_plain:
                escaped_teor = teor_plain[:600].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")
                teor_flow = f"<br/><font color='#526572' size='6.5'><b>Teor:</b> {escaped_teor}</font>"
            qdata.append([
                Paragraph(f"<b>{q['question_code']}</b><br/>{q['question_title'][:200]}"
                          f"<br/>Item: {q['item_name'][:120]}", styles["Normal"]),
                Paragraph(f"Resp: {(q['resposta'] or '')[:500]}<br/>Fund: {fund_link}{teor_flow}"
                          f"<br/>Aval: {q['status']} ({q['score']}/{q['max_display']})"
                          f"<br/>Evid: {(q['evidence_text'] or '')[:500]}"
                          f"<br/>Docs: {docs[:300]}", styles["Normal"]),
            ])
        if qdata:
            t = Table([["Pergunta", "Detalhe"]] + qdata, colWidths=[55 * mm, 115 * mm], repeatRows=1)
            t.setStyle(_table_style())
            story.append(t)
    story.extend(_methodology_flow(styles))
    doc.build(story)
    return path


def _fmt(v) -> str:
    if v is None:
        return "—"
    try:
        f = float(v)
        return f"{f:g}"
    except (TypeError, ValueError):
        return str(v)


def generate_general_xlsx(rows: list[dict], cards: dict) -> Path:
    EXPORTACOES_DIR.mkdir(parents=True, exist_ok=True)
    path = EXPORTACOES_DIR / f"relatorio_geral_{_stamp()}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Métrica", "Valor"])
    for k, v in cards.items():
        ws.append([k, v])
    ws2 = wb.create_sheet("Resultados por unidade")
    ws2.append(["UF", "Unidade"]
               + [DIMENSION_NAMES[sheet] for sheet in BASE_WEIGHTS]
               + ["Nota-base", "Bônus", "Nota final", "Classificação"])
    for r in rows:
        d = r["dim"]
        ws2.append([r["uf"], r["unidade_label"]]
                   + [d.get(sheet) for sheet in BASE_WEIGHTS]
                   + [
                    r["base_score"], r["bonus_applied"], r["final_score"],
                    r["classification"]])
    ws3 = wb.create_sheet("Metodologia resumida")
    ws3.append(["Nota", NOTA_METODOLOGICA])
    for wsx in (ws, ws2, ws3):
        for cell in wsx[1]:
            cell.font = Font(bold=True)
    wb.save(str(path))
    return path


def generate_unit_xlsx(entity_key: str) -> Path:
    detail = read_unit_detail(entity_key)
    ent = detail["entity"]
    res = detail["result"]
    EXPORTACOES_DIR.mkdir(parents=True, exist_ok=True)
    path = EXPORTACOES_DIR / f"relatorio_{entity_key}_{_stamp()}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Campo", "Valor"])
    for k in ("situacao", "base_score", "bonus_available", "bonus_applied",
              "final_score", "meets_minimum_parameters", "classification"):
        ws.append([k, res.get(k)])
    ws.append(["global_final_minimum", GLOBAL_BASE_MINIMUM])
    for sheet in BASE_WEIGHTS:
        ws.append([f"minimo_{sheet}", DIMENSION_MINIMUMS.get(sheet, "")])
    ws.append(["uf", ent["uf"]])
    ws.append(["unidade", ent["unidade_label"]])
    ws.append(["nota_metodologica", NOTA_METODOLOGICA])
    ws2 = wb.create_sheet("Avaliação detalhada")
    ws2.append(["Dimensão", "Código", "Pergunta", "Item", "Resposta do diagnóstico",
                "Fundamentação", "Avaliação", "Pontuação", "Evidência"])
    for dim in detail["dimensions"]:
        for q in dim["questions"]:
            ws2.append([dim["dimension_name"], q["question_code"], q["question_title"], q["item_name"],
                        q["resposta"], q["fundamentacao"], q["status"],
                        q["score"], q["evidence_text"]])
    ws3 = wb.create_sheet("Anexos")
    ws3.append(["Pergunta", "Arquivo", "Tamanho", "SHA-256", "Ativo"])
    for dim in detail["dimensions"]:
        for q in dim["questions"]:
            for a in q["attachments"]:
                ws3.append([q["occurrence_key"], a["original_filename"],
                            a["size_bytes"], a["sha256"], "TRUE"])
    for c in ws2.columns:
        pass
    for wsx in (ws, ws2, ws3):
        wsx.sheet_view.showGridLines = True
        for cell in wsx[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(str(path))
    return path
