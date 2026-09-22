"""Anexos: validacao, armazenamento em ANEXOS/, metadados em DB_ANEXOS."""
from __future__ import annotations

import hashlib
import mimetypes
import re
import uuid
from datetime import datetime
from pathlib import Path

import openpyxl
from filelock import FileLock, Timeout

from .config import (
    ALLOWED_EXTENSIONS,
    ANEXOS_DIR,
    BASE_DIR,
    DIMENSION_NAMES,
    MAX_UPLOAD_BYTES,
)
from .workbook_service import (
    LockedError,
    WorkbookError,
    create_backup,
    ensure_aux_sheets,
    find_question,
    parse_all_questions,
    parse_entities,
    save_atomic,
    sheet_slug,
    _acquire_lock,
    append_audit,
)

_WINDOWS_INVALID = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def sanitize_filename(name: str) -> str:
    base = Path(name or "arquivo").name  # remove path traversal
    base = _WINDOWS_INVALID.sub("_", base).strip().strip(".")
    return base[:150] or "arquivo"


def validate_file(filename: str, size: int) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise WorkbookError(
            f"Extensão '{ext or '(sem extensão)'}' não permitida. "
            f"Permitidas: {', '.join(sorted(ALLOWED_EXTENSIONS))}."
        )
    if size > MAX_UPLOAD_BYTES:
        raise WorkbookError(
            f"Arquivo excede 25 MB ({size / 1048576:.1f} MB). Reduza o tamanho e tente novamente."
        )
    return ext


def store_attachment(entity_key: str, occurrence_key: str, original_filename: str,
                     data: bytes, notes: str = "") -> dict:
    size = len(data)
    validate_file(original_filename, size)
    safe = sanitize_filename(original_filename)
    digest = hashlib.sha256(data).hexdigest()
    lock: FileLock = _acquire_lock()
    try:
        try:
            lock.acquire()
        except Timeout as exc:
            from .workbook_service import LOCKED_MSG
            raise LockedError(LOCKED_MSG) from exc
        try:
            create_backup()
            from .workbook_service import open_workbook
            wb = open_workbook()
            try:
                entities = parse_entities(wb)
                ent = next((e for e in entities if e["entity_key"] == entity_key), None)
                if ent is None:
                    raise WorkbookError(f"Unidade {entity_key} não encontrada.")
                q = find_question(wb, occurrence_key)
                all_q = parse_all_questions(wb)
                ensure_aux_sheets(wb, entities, all_q)
                folder = (BASE_DIR / "ANEXOS" / entity_key / sheet_slug(q["sheet_name"]) / q["question_code"])
                folder.mkdir(parents=True, exist_ok=True)
                stored = f"{uuid.uuid4().hex}__{safe}"
                dest = folder / stored
                dest.write_bytes(data)
                rel = dest.relative_to(BASE_DIR).as_posix()
                if not rel.startswith("ANEXOS/"):
                    raise WorkbookError(f"Caminho de anexo fora de ANEXOS: {rel}")
                mime, _ = mimetypes.guess_type(original_filename)
                ws = wb["DB_ANEXOS"]
                try:
                    last = ws.cell(ws.max_row, 1).value
                    nid = int(last) + 1 if ws.max_row > 1 and last is not None else 1
                except (TypeError, ValueError):
                    nid = ws.max_row
                uploaded_at = datetime.now().isoformat(timespec="seconds")
                ws.append([
                    nid, entity_key, ent["uf"], ent["unidade_label"],
                    q["sheet_name"], DIMENSION_NAMES.get(q["sheet_name"], q["sheet_name"]),
                    q["item_name"], q["question_code"], occurrence_key,
                    original_filename, stored, rel, mime or "application/octet-stream",
                    size, digest, uploaded_at, "TRUE", notes or "",
                ])
                append_audit(wb, entity_key, q["sheet_name"], q["question_code"],
                             occurrence_key, "anexo", "", original_filename,
                             "inclusao_anexo")
                save_atomic(wb)
                return {"ok": True, "id": nid, "sha256": digest, "relative_path": rel}
            finally:
                wb.close()
        finally:
            try:
                lock.release()
            except Exception:  # noqa: BLE001
                pass
    except LockedError:
        raise
    except PermissionError as exc:
        from .workbook_service import LOCKED_MSG
        raise LockedError(LOCKED_MSG) from exc


def unlink_attachment(attachment_id) -> dict:
    """Desvincula (active=FALSE); nunca apaga o arquivo fisico."""
    lock = _acquire_lock()
    try:
        try:
            lock.acquire()
        except Timeout as exc:
            from .workbook_service import LOCKED_MSG
            raise LockedError(LOCKED_MSG) from exc
        try:
            create_backup()
            from .workbook_service import open_workbook
            wb = open_workbook()
            try:
                if "DB_ANEXOS" not in wb.sheetnames:
                    raise WorkbookError("Aba DB_ANEXOS não encontrada.")
                ws = wb["DB_ANEXOS"]
                found = None
                for r in range(2, ws.max_row + 1):
                    if str(ws.cell(r, 1).value) == str(attachment_id):
                        found = r
                        break
                if found is None:
                    raise WorkbookError(f"Anexo {attachment_id} não encontrado.")
                ws.cell(found, 17).value = "FALSE"
                append_audit(
                    wb, str(ws.cell(found, 2).value), str(ws.cell(found, 5).value),
                    str(ws.cell(found, 8).value), str(ws.cell(found, 9).value),
                    "anexo", str(ws.cell(found, 10).value), "desvinculado",
                    "desvinculacao_anexo",
                )
                save_atomic(wb)
                return {"ok": True}
            finally:
                wb.close()
        finally:
            try:
                lock.release()
            except Exception:  # noqa: BLE001
                pass
    except LockedError:
        raise
    except PermissionError as exc:
        from .workbook_service import LOCKED_MSG
        raise LockedError(LOCKED_MSG) from exc


def get_attachment(attachment_id) -> dict:
    """Retorna metadados + path do arquivo (somente id valido, sem path do usuario)."""
    from .workbook_service import open_workbook
    wb = open_workbook()
    try:
        if "DB_ANEXOS" not in wb.sheetnames:
            raise WorkbookError("Aba DB_ANEXOS não encontrada.")
        ws = wb["DB_ANEXOS"]
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(r, 1).value) == str(attachment_id):
                active = str(ws.cell(r, 17).value).strip().upper()
                if active not in ("TRUE", "VERDADEIRO", "1", "SIM", "YES", "Y"):
                    raise WorkbookError("Anexo desvinculado.")
                rel = str(ws.cell(r, 12).value or "")
                if not rel or ".." in rel or rel.startswith(("/", "\\")):
                    raise WorkbookError("Caminho de anexo inválido.")
                # resolve dentro do projeto; impede escape
                full = (BASE_DIR / Path(rel)).resolve()
                base = BASE_DIR.resolve()
                if full != base and base not in full.parents:
                    raise WorkbookError("Caminho de anexo fora do workspace.")
                if not full.exists():
                    raise WorkbookError("Arquivo físico do anexo não encontrado.")
                return {
                    "original_filename": str(ws.cell(r, 10).value or "anexo"),
                    "mime_type": str(ws.cell(r, 13).value or "application/octet-stream"),
                    "path": full,
                }
        raise WorkbookError(f"Anexo {attachment_id} não encontrado.")
    finally:
        wb.close()
